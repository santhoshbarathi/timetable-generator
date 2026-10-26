"""
Intelligent Department Timetable Scheduler - Flask Application

Stack: Flask (Python), SQLite, Vanilla HTML/CSS/JS.
Provides RESTful APIs with server-side validation, error handling,
automatic timetable regeneration on data mutations, and static web serving.
"""

import sqlite3
from typing import Dict, Any, Tuple
import csv
import io
from flask import Flask, request, jsonify, render_template, Response

from config import (
    DEFAULT_PART_TIME_WEEKLY_PERIODS,
    MAX_FULL_TIME_WEEKLY_PERIODS,
    WORK_TYPE_FULL_TIME,
    WORK_TYPE_PART_TIME,
    PART_TIME_ALLOWED_PERIODS,
    FULL_TIME_ALLOWED_PERIODS,
    PERIOD_TIMINGS,
    BREAK_TIMINGS,
    LUNCH_BREAK_START,
    LUNCH_BREAK_END,
    DAYS_OF_WEEK,
    TOTAL_PERIODS_PER_DAY,
    MIN_PERIODS_PER_WEEK,
    MAX_PERIODS_PER_WEEK,
)
from database import init_db, seed_sample_data_if_empty, get_db_connection
from scheduler import TimetableScheduler, SchedulingError
from supabase import create_client, Client

# Supabase Project Configuration & Client Setup
SUPABASE_URL = "https://cxzeoddfvqomshlaegla.supabase.co"
SUPABASE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImN4emVvZGRmdnFvbXNobGFlZ2xhIiwicm9sZSI6ImFub24iLCJpYXQiOjE3OTMwMjg2MDEsImV4cCI6MjEwODYwNDYwMX0.3k5lcxEDNXzCF9wo5lGSvNafwwreeNGRZB_qDG98_YQ"

supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)

app = Flask(__name__)


# ---------------------------------------------------------------------------
# Initialization Hooks
# ---------------------------------------------------------------------------
@app.before_request
def setup_if_needed():
    """Ensure database exists and seeds sample data on the first request."""
    if not hasattr(app, "_db_initialized"):
        init_db()
        seeded = seed_sample_data_if_empty()
        if seeded:
            try:
                TimetableScheduler().generate()
            except Exception:
                pass  # Seeded timetable will be regenerated on demand if needed
        app._db_initialized = True


# ---------------------------------------------------------------------------
# Helper: Safe Auto-Regeneration
# ---------------------------------------------------------------------------
def try_auto_regenerate() -> Tuple[bool, str]:
    """
    Attempts to automatically regenerate the timetable after data modifications.
    Returns (success, message).
    """
    try:
        scheduler = TimetableScheduler()
        res = scheduler.generate()
        return True, "Timetable regenerated successfully."
    except SchedulingError as se:
        return False, f"Data saved, but timetable regeneration note: {se.message}"
    except Exception as e:
        return False, "Data saved, but automatic timetable regeneration encountered an issue."


# ---------------------------------------------------------------------------
# Web Page Route
# ---------------------------------------------------------------------------
@app.route("/")
def index():
    """Serves the main application dashboard."""
    return render_template("index.html")


# ---------------------------------------------------------------------------
# Authentication Routes (Supabase)
# ---------------------------------------------------------------------------
@app.route("/api/signup", methods=["POST"])
def signup():
    """Registers a new user via Supabase Auth."""
    data = request.get_json() or {}
    email = (data.get("email") or "").strip()
    password = data.get("password") or ""

    if not email or not password:
        return jsonify({"error": "Email and password are required."}), 400
    if len(password) < 6:
        return jsonify({"error": "Password must be at least 6 characters."}), 400

    try:
        res = supabase.auth.sign_up({"email": email, "password": password})
        return jsonify({
            "success": True,
            "message": "Account created! Please check your email to verify, then log in.",
            "user_id": res.user.id if res.user else None
        }), 201
    except Exception as e:
        err_msg = str(e)
        if "already registered" in err_msg.lower() or "already exists" in err_msg.lower():
            return jsonify({"error": "This email is already registered. Please log in instead."}), 409
        return jsonify({"error": "Sign-up failed. Please try again."}), 400


@app.route("/api/login", methods=["POST"])
def login():
    """Authenticates a user via Supabase Auth and returns a session token."""
    data = request.get_json() or {}
    email = (data.get("email") or "").strip()
    password = data.get("password") or ""

    if not email or not password:
        return jsonify({"error": "Email and password are required."}), 400

    try:
        res = supabase.auth.sign_in_with_password({"email": email, "password": password})
        return jsonify({
            "success": True,
            "token": res.session.access_token,
            "email": res.user.email,
            "user_id": res.user.id
        })
    except Exception as e:
        return jsonify({"error": "Invalid email or password. Please try again."}), 401


@app.route("/api/logout", methods=["POST"])
def logout():
    """Signs the user out via Supabase Auth."""
    try:
        supabase.auth.sign_out()
    except Exception:
        pass
    return jsonify({"success": True, "message": "Logged out successfully."})


# ---------------------------------------------------------------------------
# Configuration API
# ---------------------------------------------------------------------------
@app.route("/api/config", methods=["GET"])
def get_config():
    """Returns timing constants, break definitions, and scheduling parameters."""
    return jsonify({
        "college_window": "09:00 AM - 04:45 PM",
        "total_periods": TOTAL_PERIODS_PER_DAY,
        "days": DAYS_OF_WEEK,
        "lunch_break": {
            "start": LUNCH_BREAK_START,
            "end": LUNCH_BREAK_END,
            "after_period": 4,
            "note": "Fixed lunch break as defined in project requirements"
        },
        "short_breaks": {
            "break_1": {"after_period": 2, "start": "10:40", "end": "10:50", "note": "Automatic placement between Period 2 & 3"},
            "break_2": {"after_period": 6, "start": "14:45", "end": "15:05", "note": "Automatic placement between Period 6 & 7"}
        },
        "part_time_window": {
            "start": "09:00",
            "end": "12:30",
            "periods": PART_TIME_ALLOWED_PERIODS,
            "default_weekly_limit": DEFAULT_PART_TIME_WEEKLY_PERIODS
        },
        "full_time_window": {
            "start": "09:00",
            "end": "16:45",
            "periods": FULL_TIME_ALLOWED_PERIODS,
            "default_weekly_limit": MAX_FULL_TIME_WEEKLY_PERIODS
        },
        "period_timings": PERIOD_TIMINGS,
        "break_timings": BREAK_TIMINGS,
        "compulsory_faculty_hours": 16,
        "zero_free_periods_policy": True
    })


# ---------------------------------------------------------------------------
# Classes / Sections APIs
# ---------------------------------------------------------------------------
@app.route("/api/classes", methods=["GET"])
def list_classes():
    """Returns all classes with their enrolled subjects and period counts."""
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT c.*, 
               COUNT(s.id) AS subject_count,
               COALESCE(SUM(s.periods_per_week), 0) AS total_weekly_periods
        FROM classes c
        LEFT JOIN subjects s ON c.id = s.class_id
        GROUP BY c.id
        ORDER BY c.name;
    """)
    rows = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return jsonify(rows)


@app.route("/api/classes", methods=["POST"])
def add_class():
    """Adds a new class and triggers automatic timetable regeneration."""
    data = request.get_json() or {}
    name = (data.get("name") or "").strip()
    code = (data.get("code") or "").strip().upper()
    try:
        semester = int(data.get("semester", 1))
        if semester < 1 or semester > 8:
            semester = 1
    except (ValueError, TypeError):
        semester = 1

    if not name or len(name) < 2:
        return jsonify({"error": "Class name is required (minimum 2 characters)."}), 400
    if not code or len(code) < 2:
        return jsonify({"error": "Class code is required (e.g. 'CSE-4A')."}), 400

    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("INSERT INTO classes (name, code, semester) VALUES (?, ?, ?);", (name, code, semester))
        new_id = cursor.lastrowid
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": "A class with this name or code already exists."}), 400
    conn.close()

    regen_ok, regen_msg = try_auto_regenerate()
    return jsonify({
        "success": True,
        "id": new_id,
        "message": f"Class '{name}' created successfully.",
        "regeneration": {"success": regen_ok, "message": regen_msg}
    }), 201


@app.route("/api/classes/<int:class_id>", methods=["DELETE"])
def delete_class(class_id: int):
    """Deletes a class and cascade-deletes subjects and timetable entries."""
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT name FROM classes WHERE id = ?;", (class_id,))
    row = cursor.fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "Class not found."}), 404

    class_name = row["name"]
    cursor.execute("DELETE FROM classes WHERE id = ?;", (class_id,))
    conn.commit()
    conn.close()

    regen_ok, regen_msg = try_auto_regenerate()
    return jsonify({
        "success": True,
        "message": f"Class '{class_name}' deleted.",
        "regeneration": {"success": regen_ok, "message": regen_msg}
    })


# ---------------------------------------------------------------------------
# Faculty Management APIs
# ---------------------------------------------------------------------------
@app.route("/api/faculty", methods=["GET"])
def list_faculty():
    """Returns all faculty with active timetable assigned period counts."""
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT f.*,
               COUNT(t.id) AS assigned_periods
        FROM faculty f
        LEFT JOIN timetable_entries t ON f.id = t.faculty_id
        GROUP BY f.id
        ORDER BY f.name;
    """)
    faculty_list = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return jsonify(faculty_list)


@app.route("/api/faculty", methods=["POST"])
def add_faculty():
    """
    Adds a new faculty member with validation and triggers automatic regeneration.
    Part-time faculty default to DEFAULT_PART_TIME_WEEKLY_PERIODS unless specified.
    """
    data = request.get_json() or {}
    name = (data.get("name") or "").strip()
    department = (data.get("department") or "").strip()
    work_type = (data.get("work_type") or "").strip().lower()

    if not name:
        return jsonify({"error": "Faculty name is required."}), 400
    if work_type not in (WORK_TYPE_FULL_TIME, WORK_TYPE_PART_TIME):
        return jsonify({"error": "Work type must be either 'full_time' or 'part_time'."}), 400

    # Determine weekly period limit
    limit_raw = data.get("weekly_period_limit")
    if limit_raw is not None and str(limit_raw).isdigit():
        weekly_limit = int(limit_raw)
    else:
        weekly_limit = (
            DEFAULT_PART_TIME_WEEKLY_PERIODS
            if work_type == WORK_TYPE_PART_TIME
            else MAX_FULL_TIME_WEEKLY_PERIODS
        )

    if weekly_limit < 1 or weekly_limit > 40:
        return jsonify({"error": "Weekly period limit must be between 1 and 40."}), 400

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO faculty (name, department, work_type, weekly_period_limit)
        VALUES (?, ?, ?, ?);
    """, (name, department, work_type, weekly_limit))
    new_id = cursor.lastrowid
    conn.commit()
    conn.close()

    regen_ok, regen_msg = try_auto_regenerate()
    return jsonify({
        "success": True,
        "id": new_id,
        "message": f"Faculty '{name}' added successfully.",
        "regeneration": {"success": regen_ok, "message": regen_msg}
    }), 201


@app.route("/api/faculty/<int:faculty_id>", methods=["PUT"])
def update_faculty(faculty_id: int):
    """Updates an existing faculty member and triggers timetable regeneration."""
    data = request.get_json() or {}
    name = (data.get("name") or "").strip()
    department = (data.get("department") or "").strip()
    work_type = (data.get("work_type") or "").strip().lower()

    if not name:
        return jsonify({"error": "Faculty name is required."}), 400
    if work_type not in (WORK_TYPE_FULL_TIME, WORK_TYPE_PART_TIME):
        return jsonify({"error": "Work type must be either 'full_time' or 'part_time'."}), 400

    limit_raw = data.get("weekly_period_limit")
    if limit_raw is not None and str(limit_raw).isdigit():
        weekly_limit = int(limit_raw)
    else:
        weekly_limit = (
            DEFAULT_PART_TIME_WEEKLY_PERIODS
            if work_type == WORK_TYPE_PART_TIME
            else MAX_FULL_TIME_WEEKLY_PERIODS
        )

    if weekly_limit < 1 or weekly_limit > 40:
        return jsonify({"error": "Weekly period limit must be between 1 and 40."}), 400

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id FROM faculty WHERE id = ?;", (faculty_id,))
    if not cursor.fetchone():
        conn.close()
        return jsonify({"error": "Faculty member not found."}), 404

    cursor.execute("""
        UPDATE faculty
        SET name = ?, department = ?, work_type = ?, weekly_period_limit = ?
        WHERE id = ?;
    """, (name, department, work_type, weekly_limit, faculty_id))
    conn.commit()
    conn.close()

    regen_ok, regen_msg = try_auto_regenerate()
    return jsonify({
        "success": True,
        "message": f"Faculty '{name}' updated successfully.",
        "regeneration": {"success": regen_ok, "message": regen_msg}
    })


@app.route("/api/faculty/<int:faculty_id>", methods=["DELETE"])
def delete_faculty(faculty_id: int):
    """Deletes a faculty member and triggers automatic regeneration."""
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT name FROM faculty WHERE id = ?;", (faculty_id,))
    row = cursor.fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "Faculty member not found."}), 404

    faculty_name = row["name"]
    cursor.execute("DELETE FROM faculty WHERE id = ?;", (faculty_id,))
    conn.commit()
    conn.close()

    regen_ok, regen_msg = try_auto_regenerate()
    return jsonify({
        "success": True,
        "message": f"Faculty '{faculty_name}' removed.",
        "regeneration": {"success": regen_ok, "message": regen_msg}
    })


# ---------------------------------------------------------------------------
# Subject Management APIs
# ---------------------------------------------------------------------------
@app.route("/api/subjects", methods=["GET"])
def list_subjects():
    """Returns subjects, optionally filtered by class_id."""
    class_id = request.args.get("class_id")
    conn = get_db_connection()
    cursor = conn.cursor()

    base_query = """
        SELECT s.*, 
               c.name AS class_name, c.code AS class_code,
               f.name AS faculty_name, f.work_type AS faculty_work_type
        FROM subjects s
        JOIN classes c ON s.class_id = c.id
        LEFT JOIN faculty f ON s.faculty_id = f.id
    """
    params = []
    if class_id and class_id.isdigit():
        base_query += " WHERE s.class_id = ?"
        params.append(int(class_id))

    base_query += " ORDER BY c.name, s.name;"
    cursor.execute(base_query, params)
    subjects = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return jsonify(subjects)


@app.route("/api/subjects", methods=["POST"])
def add_subject():
    """
    Creates a new subject for a class and triggers automatic regeneration.
    Enforces periods_per_week between 1 and 5 (max 1 per day constraint).
    """
    data = request.get_json() or {}
    class_id = data.get("class_id")
    name = (data.get("name") or "").strip()
    code = (data.get("code") or "").strip().upper()
    periods_raw = data.get("periods_per_week")
    faculty_id = data.get("faculty_id")

    if not class_id:
        return jsonify({"error": "Class selection is required."}), 400
    if not name:
        return jsonify({"error": "Subject name is required."}), 400
    if not code:
        return jsonify({"error": "Subject code is required (e.g. 'CS201')."}), 400

    try:
        periods_per_week = int(periods_raw)
        if periods_per_week < MIN_PERIODS_PER_WEEK or periods_per_week > MAX_PERIODS_PER_WEEK:
            return jsonify({
                "error": f"Periods per week must be between {MIN_PERIODS_PER_WEEK} and {MAX_PERIODS_PER_WEEK}."
            }), 400
    except (ValueError, TypeError):
        return jsonify({"error": "Invalid periods per week value."}), 400

    # Optional faculty_id validation
    f_id_clean = None
    if faculty_id not in (None, "", "null", 0):
        try:
            f_id_clean = int(faculty_id)
        except ValueError:
            return jsonify({"error": "Invalid faculty ID."}), 400

    conn = get_db_connection()
    cursor = conn.cursor()

    # Verify class exists
    cursor.execute("SELECT id FROM classes WHERE id = ?;", (class_id,))
    if not cursor.fetchone():
        conn.close()
        return jsonify({"error": "Specified class does not exist."}), 400

    # Verify faculty exists if provided
    if f_id_clean:
        cursor.execute("SELECT id FROM faculty WHERE id = ?;", (f_id_clean,))
        if not cursor.fetchone():
            conn.close()
            return jsonify({"error": "Specified faculty does not exist."}), 400

    try:
        cursor.execute("""
            INSERT INTO subjects (class_id, name, code, periods_per_week, faculty_id)
            VALUES (?, ?, ?, ?, ?);
        """, (class_id, name, code, periods_per_week, f_id_clean))
        new_id = cursor.lastrowid
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": "A subject with this code already exists in this class."}), 400
    conn.close()

    regen_ok, regen_msg = try_auto_regenerate()
    return jsonify({
        "success": True,
        "id": new_id,
        "message": f"Subject '{name}' added successfully.",
        "regeneration": {"success": regen_ok, "message": regen_msg}
    }), 201


@app.route("/api/subjects/<int:subject_id>", methods=["PUT"])
def update_subject(subject_id: int):
    """Updates subject details and triggers automatic timetable regeneration."""
    data = request.get_json() or {}
    name = (data.get("name") or "").strip()
    code = (data.get("code") or "").strip().upper()
    periods_raw = data.get("periods_per_week")
    faculty_id = data.get("faculty_id")

    if not name:
        return jsonify({"error": "Subject name is required."}), 400
    if not code:
        return jsonify({"error": "Subject code is required."}), 400

    try:
        periods_per_week = int(periods_raw)
        if periods_per_week < MIN_PERIODS_PER_WEEK or periods_per_week > MAX_PERIODS_PER_WEEK:
            return jsonify({
                "error": f"Periods per week must be between {MIN_PERIODS_PER_WEEK} and {MAX_PERIODS_PER_WEEK}."
            }), 400
    except (ValueError, TypeError):
        return jsonify({"error": "Invalid periods per week value."}), 400

    f_id_clean = None
    if faculty_id not in (None, "", "null", 0):
        try:
            f_id_clean = int(faculty_id)
        except ValueError:
            return jsonify({"error": "Invalid faculty ID."}), 400

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, class_id FROM subjects WHERE id = ?;", (subject_id,))
    row = cursor.fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "Subject not found."}), 404

    class_id = row["class_id"]

    try:
        cursor.execute("""
            UPDATE subjects
            SET name = ?, code = ?, periods_per_week = ?, faculty_id = ?
            WHERE id = ?;
        """, (name, code, periods_per_week, f_id_clean, subject_id))
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": "Another subject with this code already exists in this class."}), 400
    conn.close()

    regen_ok, regen_msg = try_auto_regenerate()
    return jsonify({
        "success": True,
        "message": f"Subject '{name}' updated successfully.",
        "regeneration": {"success": regen_ok, "message": regen_msg}
    })


@app.route("/api/subjects/<int:subject_id>", methods=["DELETE"])
def delete_subject(subject_id: int):
    """Deletes a subject and triggers automatic timetable regeneration."""
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT name FROM subjects WHERE id = ?;", (subject_id,))
    row = cursor.fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "Subject not found."}), 404

    subject_name = row["name"]
    cursor.execute("DELETE FROM subjects WHERE id = ?;", (subject_id,))
    conn.commit()
    conn.close()

    regen_ok, regen_msg = try_auto_regenerate()
    return jsonify({
        "success": True,
        "message": f"Subject '{subject_name}' deleted.",
        "regeneration": {"success": regen_ok, "message": regen_msg}
    })


# ---------------------------------------------------------------------------
# Timetable APIs
# ---------------------------------------------------------------------------
@app.route("/api/timetable", methods=["GET"])
def get_timetable():
    """
    Fetches the current timetable structured by class and day,
    along with break markers and faculty utilization statistics.
    """
    conn = get_db_connection()
    cursor = conn.cursor()

    # Load classes
    cursor.execute("SELECT * FROM classes ORDER BY id;")
    classes = [dict(r) for r in cursor.fetchall()]

    # Load timetable entries
    cursor.execute("""
        SELECT t.*, 
               s.name AS subject_name, s.code AS subject_code,
               f.name AS faculty_name, f.work_type AS faculty_work_type
        FROM timetable_entries t
        JOIN subjects s ON t.subject_id = s.id
        JOIN faculty f ON t.faculty_id = f.id
        ORDER BY t.class_id, t.day, t.period_num;
    """)
    entries = [dict(r) for r in cursor.fetchall()]

    # Load faculty stats
    cursor.execute("""
        SELECT f.id, f.name, f.work_type, f.weekly_period_limit,
               COUNT(t.id) AS assigned_periods
        FROM faculty f
        LEFT JOIN timetable_entries t ON f.id = t.faculty_id
        GROUP BY f.id
        ORDER BY f.name;
    """)
    faculty_stats = [dict(r) for r in cursor.fetchall()]
    conn.close()

    schedules_by_class: Dict[int, Dict[str, Any]] = {}
    for c in classes:
        schedules_by_class[c["id"]] = {
            "class_id": c["id"],
            "class_name": c["name"],
            "class_code": c["code"],
            "semester": c.get("semester", 1),
            "grid": {day: {} for day in DAYS_OF_WEEK},
        }

    for e in entries:
        cid = e["class_id"]
        day = e["day"]
        p = e["period_num"]
        if cid in schedules_by_class and day in schedules_by_class[cid]["grid"]:
            schedules_by_class[cid]["grid"][day][p] = {
                "id": e["id"],
                "entry_id": e["id"],
                "period_num": p,
                "timing": PERIOD_TIMINGS.get(p, {}),
                "subject_id": e["subject_id"],
                "subject_name": e["subject_name"],
                "subject_code": e["subject_code"],
                "faculty_id": e["faculty_id"],
                "faculty_name": e["faculty_name"],
                "faculty_work_type": e["faculty_work_type"],
            }

    return jsonify({
        "total_classes": len(classes),
        "total_entries": len(entries),
        "schedules_by_class": schedules_by_class,
        "faculty_utilization": faculty_stats,
        "period_timings": PERIOD_TIMINGS,
        "break_timings": BREAK_TIMINGS,
        "days_of_week": DAYS_OF_WEEK,
    })


@app.route("/api/timetable/swap", methods=["POST"])
def swap_timetable_slot():
    """
    Interactive Timetable Slot Swapping API with Live Pre-Flight Conflict Checking.
    Supports:
      1. Swapping source entry with an empty target slot:
         { "source_entry_id": int, "target_day": str, "target_period": int }
      2. Swapping two distinct entries:
         { "source_id": int, "target_id": int }
      3. Also accepts { "source_entry_id": int, "target_entry_id": int }
    Validates:
      - Constraint 1 (Faculty Double-Booking): Source faculty is not already booked in target (day, period).
      - Constraint 2 (Part-Time Faculty Window): Part-time faculty cannot teach Period > 4 (after 12:30 PM).
      - Reciprocal constraints if target slot is also occupied.
    Updates both Supabase and local SQLite to maintain synchronization.
    """
    data = request.get_json() or {}

    source_entry_id = data.get("source_entry_id") or data.get("source_id")
    target_entry_id = data.get("target_entry_id") or data.get("target_id")
    target_day = data.get("target_day")
    target_period = data.get("target_period")

    if not source_entry_id:
        return jsonify({"success": False, "conflict": True, "message": "Source slot identifier is required."}), 400

    conn = get_db_connection()
    cursor = conn.cursor()

    # Load source entry with faculty, subject, and class details
    cursor.execute("""
        SELECT t.*, f.name AS faculty_name, f.work_type AS faculty_work_type,
               s.name AS subject_name, s.code AS subject_code,
               c.name AS class_name, c.code AS class_code
        FROM timetable_entries t
        JOIN faculty f ON t.faculty_id = f.id
        JOIN subjects s ON t.subject_id = s.id
        JOIN classes c ON t.class_id = c.id
        WHERE t.id = ?;
    """, (source_entry_id,))
    source_row = cursor.fetchone()
    if not source_row:
        conn.close()
        return jsonify({"success": False, "conflict": True, "message": "Source timetable entry not found."}), 404
    source_entry = dict(source_row)

    # Determine target entry if target_entry_id provided or target_day/period occupied
    target_entry = None
    if target_entry_id:
        cursor.execute("""
            SELECT t.*, f.name AS faculty_name, f.work_type AS faculty_work_type,
                   s.name AS subject_name, s.code AS subject_code,
                   c.name AS class_name, c.code AS class_code
            FROM timetable_entries t
            JOIN faculty f ON t.faculty_id = f.id
            JOIN subjects s ON t.subject_id = s.id
            JOIN classes c ON t.class_id = c.id
            WHERE t.id = ?;
        """, (target_entry_id,))
        t_row = cursor.fetchone()
        if t_row:
            target_entry = dict(t_row)
            target_day = target_entry["day"]
            target_period = target_entry["period_num"]
    elif target_day and target_period:
        target_period = int(target_period)
        # Check if the class already has an entry at target_day, target_period
        cursor.execute("""
            SELECT t.*, f.name AS faculty_name, f.work_type AS faculty_work_type,
                   s.name AS subject_name, s.code AS subject_code,
                   c.name AS class_name, c.code AS class_code
            FROM timetable_entries t
            JOIN faculty f ON t.faculty_id = f.id
            JOIN subjects s ON t.subject_id = s.id
            JOIN classes c ON t.class_id = c.id
            WHERE t.class_id = ? AND t.day = ? AND t.period_num = ?;
        """, (source_entry["class_id"], target_day, target_period))
        t_row = cursor.fetchone()
        if t_row:
            target_entry = dict(t_row)
    else:
        conn.close()
        return jsonify({"success": False, "conflict": True, "message": "Target day and period or target slot ID required."}), 400

    target_period = int(target_period)

    # Validate target_period is within valid teaching periods 1-8
    if target_period < 1 or target_period > TOTAL_PERIODS_PER_DAY:
        conn.close()
        return jsonify({"success": False, "conflict": True, "message": f"Period {target_period} is outside operating periods 1-8."}), 422

    # If source and target are the same slot, no-op
    if target_entry and target_entry["id"] == source_entry["id"]:
        conn.close()
        return jsonify({"success": True, "message": "Source and target slots are identical."}), 200

    # -------------------------------------------------------------------------
    # PRE-FLIGHT CONSTRAINT 1: Check Source Faculty on Target (Day, Period)
    # -------------------------------------------------------------------------
    # Check if source faculty has part_time window limit
    if source_entry["faculty_work_type"] == WORK_TYPE_PART_TIME and target_period > 4:
        conn.close()
        return jsonify({
            "success": False,
            "conflict": True,
            "message": f"Conflict: Prof. {source_entry['faculty_name']} is part-time (restricted to Periods 1–4, 09:00 AM – 12:30 PM only). Cannot teach Period {target_period}!"
        }), 422

    # Check faculty double booking in Supabase / SQLite
    # Find any booking for source_entry['faculty_id'] on target_day, target_period excluding target_entry (which will move)
    cursor.execute("""
        SELECT t.id, c.code AS class_code, c.name AS class_name
        FROM timetable_entries t
        JOIN classes c ON t.class_id = c.id
        WHERE t.faculty_id = ? AND t.day = ? AND t.period_num = ?
          AND t.id != ? AND t.id != ?;
    """, (
        source_entry["faculty_id"],
        target_day,
        target_period,
        source_entry["id"],
        target_entry["id"] if target_entry else -1
    ))
    fac_conflict = cursor.fetchone()
    if fac_conflict:
        conn.close()
        return jsonify({
            "success": False,
            "conflict": True,
            "message": f"Conflict: Prof. {source_entry['faculty_name']} is already teaching {fac_conflict['class_code']} on {target_day} Period {target_period}!"
        }), 422

    # Also verify live against Supabase if configured
    try:
        sb_query = supabase.table("timetable_entries").select("id, class_id")\
            .eq("faculty_id", source_entry["faculty_id"])\
            .eq("day", target_day)\
            .eq("period_num", target_period)\
            .neq("id", source_entry["id"])
        if target_entry:
            sb_query = sb_query.neq("id", target_entry["id"])
        sb_res = sb_query.execute()
        if sb_res.data and len(sb_res.data) > 0:
            other_cid = sb_res.data[0].get("class_id")
            c_row = supabase.table("classes").select("code").eq("id", other_cid).execute()
            other_code = c_row.data[0]["code"] if c_row.data else f"Class {other_cid}"
            conn.close()
            return jsonify({
                "success": False,
                "conflict": True,
                "message": f"Conflict: Prof. {source_entry['faculty_name']} is already teaching {other_code} on {target_day} Period {target_period}!"
            }), 422
    except Exception:
        pass  # Supabase live check fallback to SQLite

    # -------------------------------------------------------------------------
    # PRE-FLIGHT CONSTRAINT 2: If Target slot is occupied, check Target Faculty on Source (Day, Period)
    # -------------------------------------------------------------------------
    if target_entry:
        src_day = source_entry["day"]
        src_period = int(source_entry["period_num"])

        # Check part-time limit for target faculty moving to source slot
        if target_entry["faculty_work_type"] == WORK_TYPE_PART_TIME and src_period > 4:
            conn.close()
            return jsonify({
                "success": False,
                "conflict": True,
                "message": f"Conflict: Prof. {target_entry['faculty_name']} is part-time (restricted to Periods 1–4, 09:00 AM – 12:30 PM only). Cannot teach Period {src_period}!"
            }), 422

        # Check double booking for target faculty in source slot
        cursor.execute("""
            SELECT t.id, c.code AS class_code, c.name AS class_name
            FROM timetable_entries t
            JOIN classes c ON t.class_id = c.id
            WHERE t.faculty_id = ? AND t.day = ? AND t.period_num = ?
              AND t.id != ? AND t.id != ?;
        """, (
            target_entry["faculty_id"],
            src_day,
            src_period,
            source_entry["id"],
            target_entry["id"]
        ))
        tgt_fac_conflict = cursor.fetchone()
        if tgt_fac_conflict:
            conn.close()
            return jsonify({
                "success": False,
                "conflict": True,
                "message": f"Conflict: Prof. {target_entry['faculty_name']} is already teaching {tgt_fac_conflict['class_code']} on {src_day} Period {src_period}!"
            }), 422

    # -------------------------------------------------------------------------
    # EXECUTE SWAP
    # -------------------------------------------------------------------------
    try:
        if target_entry:
            src_day, src_period = source_entry["day"], source_entry["period_num"]
            tgt_day, tgt_period = target_entry["day"], target_entry["period_num"]

            # Use a temporary period value to avoid UNIQUE(class_id, day, period_num) violation
            temp_p = 999
            cursor.execute("UPDATE timetable_entries SET period_num = ? WHERE id = ?;", (temp_p, source_entry["id"]))
            cursor.execute("UPDATE timetable_entries SET day = ?, period_num = ? WHERE id = ?;", (src_day, src_period, target_entry["id"]))
            cursor.execute("UPDATE timetable_entries SET day = ?, period_num = ? WHERE id = ?;", (tgt_day, tgt_period, source_entry["id"]))
            conn.commit()

            # Sync to Supabase
            try:
                supabase.table("timetable_entries").update({"day": tgt_day, "period_num": tgt_period}).eq("id", source_entry["id"]).execute()
                supabase.table("timetable_entries").update({"day": src_day, "period_num": src_period}).eq("id", target_entry["id"]).execute()
            except Exception:
                pass

            msg = f"Swapped '{source_entry['subject_name']}' ({source_entry['day']} P{source_entry['period_num']}) with '{target_entry['subject_name']}' ({target_entry['day']} P{target_entry['period_num']}) successfully!"
        else:
            # Move source entry to empty slot
            cursor.execute("UPDATE timetable_entries SET day = ?, period_num = ? WHERE id = ?;", (target_day, target_period, source_entry["id"]))
            conn.commit()

            # Sync to Supabase
            try:
                supabase.table("timetable_entries").update({"day": target_day, "period_num": target_period}).eq("id", source_entry["id"]).execute()
            except Exception:
                pass

            msg = f"Moved '{source_entry['subject_name']}' to {target_day} Period {target_period} successfully!"

        conn.close()
        return jsonify({"success": True, "message": msg}), 200

    except Exception as ex:
        conn.rollback()
        conn.close()
        return jsonify({"success": False, "conflict": True, "message": f"Failed to execute swap: {str(ex)}"}), 500


@app.route("/api/timetable/faculty/<int:faculty_id>", methods=["GET"])
def get_faculty_timetable(faculty_id: int):
    """
    Returns the weekly timetable grid for a single faculty member,
    derived from the same timetable_entries table used by the class view.
    """
    conn = get_db_connection()
    cursor = conn.cursor()

    # Verify faculty exists
    cursor.execute("SELECT * FROM faculty WHERE id = ?;", (faculty_id,))
    fac_row = cursor.fetchone()
    if not fac_row:
        conn.close()
        return jsonify({"error": "Faculty member not found."}), 404
    fac = dict(fac_row)

    # Load this faculty's timetable entries (all classes)
    cursor.execute("""
        SELECT t.day, t.period_num,
               s.name AS subject_name, s.code AS subject_code,
               s.id   AS subject_id,
               c.name AS class_name,  c.code AS class_code
        FROM timetable_entries t
        JOIN subjects s ON t.subject_id = s.id
        JOIN classes  c ON t.class_id   = c.id
        WHERE t.faculty_id = ?
        ORDER BY t.day, t.period_num;
    """, (faculty_id,))
    entries = [dict(r) for r in cursor.fetchall()]
    conn.close()

    # Build a day → period_num → slot dict
    grid: Dict[str, Any] = {day: {} for day in DAYS_OF_WEEK}
    for e in entries:
        day = e["day"]
        p   = e["period_num"]
        if day in grid:
            grid[day][p] = {
                "period_num":   p,
                "timing":       PERIOD_TIMINGS.get(p, {}),
                "subject_id":   e["subject_id"],
                "subject_name": e["subject_name"],
                "subject_code": e["subject_code"],
                "class_name":   e["class_name"],
                "class_code":   e["class_code"],
            }

    return jsonify({
        "faculty": {
            "id":                fac["id"],
            "name":              fac["name"],
            "department":        fac["department"],
            "work_type":         fac["work_type"],
            "weekly_period_limit": fac["weekly_period_limit"],
            "assigned_periods":  len(entries),
        },
        "grid":          grid,
        "period_timings": PERIOD_TIMINGS,
        "break_timings":  BREAK_TIMINGS,
        "days_of_week":   DAYS_OF_WEEK,
    })



# ---------------------------------------------------------------------------
# Export & Download Endpoints (CSV / Excel formatted timetable downloads)
# ---------------------------------------------------------------------------
@app.route("/api/export/timetable/csv", methods=["GET"])
def export_timetable_csv():
    """
    Exports the generated timetable to a cleanly formatted CSV spreadsheet.
    Supports single class (?class_id=X) or all classes (?all=true).
    """
    class_id = request.args.get("class_id", type=int)
    export_all = request.args.get("all", "false").lower() in ("true", "1")

    conn = get_db_connection()
    cursor = conn.cursor()

    if class_id and not export_all:
        cursor.execute("SELECT * FROM classes WHERE id = ?;", (class_id,))
        classes = [dict(r) for r in cursor.fetchall()]
        if not classes:
            conn.close()
            return jsonify({"error": "Class not found."}), 404
    else:
        cursor.execute("SELECT * FROM classes ORDER BY id;")
        classes = [dict(r) for r in cursor.fetchall()]

    cursor.execute("""
        SELECT t.*, 
               s.name AS subject_name, s.code AS subject_code,
               f.name AS faculty_name
        FROM timetable_entries t
        JOIN subjects s ON t.subject_id = s.id
        JOIN faculty f ON t.faculty_id = f.id
        ORDER BY t.class_id, t.day, t.period_num;
    """)
    all_entries = [dict(r) for r in cursor.fetchall()]
    conn.close()

    entries_by_class = {}
    for e in all_entries:
        cid = e["class_id"]
        if cid not in entries_by_class:
            entries_by_class[cid] = {}
        day = e["day"]
        if day not in entries_by_class[cid]:
            entries_by_class[cid][day] = {}
        entries_by_class[cid][day][e["period_num"]] = e

    output = io.StringIO()
    output.write("\ufeff")  # UTF-8 BOM for Microsoft Excel compatibility
    writer = csv.writer(output)

    for idx, c in enumerate(classes):
        if idx > 0:
            writer.writerow([])
            writer.writerow(["=" * 80])
            writer.writerow([])

        writer.writerow(["INFO INSTITUTE OF ENGINEERING (Autonomous)"])
        writer.writerow(["Approved by AICTE, New Delhi & Affiliated to Anna University, Chennai"])
        writer.writerow(["DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING"])
        writer.writerow([f"CLASS TIME TABLE: {c['name']} ({c['code']})"])
        writer.writerow(["Academic Year: 2025-2026", "College Operating Window: 09:00 AM - 04:45 PM"])
        writer.writerow([])

        headers = [
            "Day / Time",
            f"Period 1 ({PERIOD_TIMINGS[1]['start']} - {PERIOD_TIMINGS[1]['end']})",
            f"Period 2 ({PERIOD_TIMINGS[2]['start']} - {PERIOD_TIMINGS[2]['end']})",
            "Morning Break (10:40 - 10:50)",
            f"Period 3 ({PERIOD_TIMINGS[3]['start']} - {PERIOD_TIMINGS[3]['end']})",
            f"Period 4 ({PERIOD_TIMINGS[4]['start']} - {PERIOD_TIMINGS[4]['end']})",
            "Lunch Break (12:30 - 13:05)",
            f"Period 5 ({PERIOD_TIMINGS[5]['start']} - {PERIOD_TIMINGS[5]['end']})",
            f"Period 6 ({PERIOD_TIMINGS[6]['start']} - {PERIOD_TIMINGS[6]['end']})",
            "Tea Break (14:45 - 15:05)",
            f"Period 7 ({PERIOD_TIMINGS[7]['start']} - {PERIOD_TIMINGS[7]['end']})",
            f"Period 8 ({PERIOD_TIMINGS[8]['start']} - {PERIOD_TIMINGS[8]['end']})",
        ]
        writer.writerow(headers)

        c_entries = entries_by_class.get(c["id"], {})
        for day in DAYS_OF_WEEK:
            day_grid = c_entries.get(day, {})
            row = [day]

            # P1
            e1 = day_grid.get(1)
            row.append(f"{e1['subject_name']} [{e1['subject_code']}] ({e1['faculty_name']})" if e1 else "-")
            # P2
            e2 = day_grid.get(2)
            row.append(f"{e2['subject_name']} [{e2['subject_code']}] ({e2['faculty_name']})" if e2 else "-")
            # Morning Break
            row.append("MORNING BREAK")
            # P3
            e3 = day_grid.get(3)
            row.append(f"{e3['subject_name']} [{e3['subject_code']}] ({e3['faculty_name']})" if e3 else "-")
            # P4
            e4 = day_grid.get(4)
            row.append(f"{e4['subject_name']} [{e4['subject_code']}] ({e4['faculty_name']})" if e4 else "-")
            # Lunch Break
            row.append("LUNCH BREAK")
            # P5
            e5 = day_grid.get(5)
            row.append(f"{e5['subject_name']} [{e5['subject_code']}] ({e5['faculty_name']})" if e5 else "-")
            # P6
            e6 = day_grid.get(6)
            row.append(f"{e6['subject_name']} [{e6['subject_code']}] ({e6['faculty_name']})" if e6 else "-")
            # Tea Break
            row.append("TEA BREAK")
            # P7
            e7 = day_grid.get(7)
            row.append(f"{e7['subject_name']} [{e7['subject_code']}] ({e7['faculty_name']})" if e7 else "-")
            # P8
            e8 = day_grid.get(8)
            row.append(f"{e8['subject_name']} [{e8['subject_code']}] ({e8['faculty_name']})" if e8 else "-")

            writer.writerow(row)

        writer.writerow([])
        writer.writerow(["Time Table Coordinator", "", "", "HOD / CSE", "", "", "", "Principal"])

    csv_data = output.getvalue()
    filename = f"Timetable_{classes[0]['code']}.csv" if len(classes) == 1 else "All_Class_Timetables.csv"
    return Response(
        csv_data,
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@app.route("/api/export/timetable/excel", methods=["GET"])
def export_timetable_excel():
    """
    Exports timetable to professional Microsoft Excel (.xlsx) workbook using openpyxl.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    class_id = request.args.get("class_id", type=int)
    export_all = request.args.get("all", "false").lower() in ("true", "1")

    conn = get_db_connection()
    cursor = conn.cursor()

    if class_id and not export_all:
        cursor.execute("SELECT * FROM classes WHERE id = ?;", (class_id,))
        classes = [dict(r) for r in cursor.fetchall()]
        if not classes:
            conn.close()
            return jsonify({"error": "Class not found."}), 404
    else:
        cursor.execute("SELECT * FROM classes ORDER BY id;")
        classes = [dict(r) for r in cursor.fetchall()]

    cursor.execute("""
        SELECT t.*, 
               s.name AS subject_name, s.code AS subject_code,
               f.name AS faculty_name
        FROM timetable_entries t
        JOIN subjects s ON t.subject_id = s.id
        JOIN faculty f ON t.faculty_id = f.id
        ORDER BY t.class_id, t.day, t.period_num;
    """)
    all_entries = [dict(r) for r in cursor.fetchall()]
    conn.close()

    entries_by_class = {}
    for e in all_entries:
        cid = e["class_id"]
        entries_by_class.setdefault(cid, {}).setdefault(e["day"], {})[e["period_num"]] = e

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    header_font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
    sub_font = Font(name="Arial", size=10, italic=True, color="475569")
    table_head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    day_font = Font(name="Arial", size=10, bold=True, color="0F172A")
    cell_font = Font(name="Arial", size=9)
    break_font = Font(name="Arial", size=9, bold=True, color="92400E")

    primary_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    break_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    day_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = [
        "Day / Time",
        f"Period 1\n({PERIOD_TIMINGS[1]['start']}-{PERIOD_TIMINGS[1]['end']})",
        f"Period 2\n({PERIOD_TIMINGS[2]['start']}-{PERIOD_TIMINGS[2]['end']})",
        "Morning Break\n(10:40-10:50)",
        f"Period 3\n({PERIOD_TIMINGS[3]['start']}-{PERIOD_TIMINGS[3]['end']})",
        f"Period 4\n({PERIOD_TIMINGS[4]['start']}-{PERIOD_TIMINGS[4]['end']})",
        "Lunch Break\n(12:30-13:05)",
        f"Period 5\n({PERIOD_TIMINGS[5]['start']}-{PERIOD_TIMINGS[5]['end']})",
        f"Period 6\n({PERIOD_TIMINGS[6]['start']}-{PERIOD_TIMINGS[6]['end']})",
        "Tea Break\n(14:45-15:05)",
        f"Period 7\n({PERIOD_TIMINGS[7]['start']}-{PERIOD_TIMINGS[7]['end']})",
        f"Period 8\n({PERIOD_TIMINGS[8]['start']}-{PERIOD_TIMINGS[8]['end']})",
    ]

    for c in classes:
        ws = wb.create_sheet(title=c["code"][:30])
        ws.sheet_view.showGridLines = True

        # Header Titles
        ws.merge_cells("A1:L1")
        ws["A1"] = "INFO INSTITUTE OF ENGINEERING (Autonomous)"
        ws["A1"].font = header_font
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A2:L2")
        ws["A2"] = "Approved by AICTE, New Delhi & Affiliated to Anna University, Chennai"
        ws["A2"].font = sub_font
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A3:L3")
        ws["A3"] = f"CLASS TIME TABLE: {c['name']} ({c['code']}) — Academic Year 2025-2026"
        ws["A3"].font = Font(name="Arial", size=11, bold=True, color="0F172A")
        ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 20
        ws.row_dimensions[5].height = 30

        # Table Header Row at Row 5
        for col_idx, col_name in enumerate(headers, start=1):
            cell = ws.cell(row=5, column=col_idx, value=col_name)
            cell.font = table_head_font
            cell.fill = primary_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        c_entries = entries_by_class.get(c["id"], {})
        current_row = 6

        for day in DAYS_OF_WEEK:
            ws.row_dimensions[current_row].height = 36
            day_grid = c_entries.get(day, {})

            # Day cell
            d_cell = ws.cell(row=current_row, column=1, value=day)
            d_cell.font = day_font
            d_cell.fill = day_fill
            d_cell.alignment = Alignment(horizontal="center", vertical="center")
            d_cell.border = thin_border

            # Period 1
            e1 = day_grid.get(1)
            c1 = ws.cell(row=current_row, column=2, value=f"{e1['subject_name']}\n[{e1['subject_code']}]" if e1 else "-")
            c1.font = cell_font
            c1.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c1.border = thin_border

            # Period 2
            e2 = day_grid.get(2)
            c2 = ws.cell(row=current_row, column=3, value=f"{e2['subject_name']}\n[{e2['subject_code']}]" if e2 else "-")
            c2.font = cell_font
            c2.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c2.border = thin_border

            # Morning Break
            mb = ws.cell(row=current_row, column=4, value="BREAK")
            mb.font = break_font
            mb.fill = break_fill
            mb.alignment = Alignment(horizontal="center", vertical="center")
            mb.border = thin_border

            # Period 3
            e3 = day_grid.get(3)
            c3 = ws.cell(row=current_row, column=5, value=f"{e3['subject_name']}\n[{e3['subject_code']}]" if e3 else "-")
            c3.font = cell_font
            c3.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c3.border = thin_border

            # Period 4
            e4 = day_grid.get(4)
            c4 = ws.cell(row=current_row, column=6, value=f"{e4['subject_name']}\n[{e4['subject_code']}]" if e4 else "-")
            c4.font = cell_font
            c4.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c4.border = thin_border

            # Lunch Break
            lb = ws.cell(row=current_row, column=7, value="LUNCH")
            lb.font = break_font
            lb.fill = break_fill
            lb.alignment = Alignment(horizontal="center", vertical="center")
            lb.border = thin_border

            # Period 5
            e5 = day_grid.get(5)
            c5 = ws.cell(row=current_row, column=8, value=f"{e5['subject_name']}\n[{e5['subject_code']}]" if e5 else "-")
            c5.font = cell_font
            c5.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c5.border = thin_border

            # Period 6
            e6 = day_grid.get(6)
            c6 = ws.cell(row=current_row, column=9, value=f"{e6['subject_name']}\n[{e6['subject_code']}]" if e6 else "-")
            c6.font = cell_font
            c6.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c6.border = thin_border

            # Tea Break
            tb = ws.cell(row=current_row, column=10, value="BREAK")
            tb.font = break_font
            tb.fill = break_fill
            tb.alignment = Alignment(horizontal="center", vertical="center")
            tb.border = thin_border

            # Period 7
            e7 = day_grid.get(7)
            c7 = ws.cell(row=current_row, column=11, value=f"{e7['subject_name']}\n[{e7['subject_code']}]" if e7 else "-")
            c7.font = cell_font
            c7.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c7.border = thin_border

            # Period 8
            e8 = day_grid.get(8)
            c8 = ws.cell(row=current_row, column=12, value=f"{e8['subject_name']}\n[{e8['subject_code']}]" if e8 else "-")
            c8.font = cell_font
            c8.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c8.border = thin_border

            current_row += 1

        # Signatures
        sig_row = current_row + 2
        ws.cell(row=sig_row, column=2, value="Time Table Coordinator").font = Font(name="Arial", size=10, bold=True)
        ws.cell(row=sig_row, column=6, value="HOD / CSE").font = Font(name="Arial", size=10, bold=True)
        ws.cell(row=sig_row, column=11, value="Principal").font = Font(name="Arial", size=10, bold=True)

        # Set column widths
        ws.column_dimensions["A"].width = 14
        for col_letter in ["B", "C", "E", "F", "H", "I", "K", "L"]:
            ws.column_dimensions[col_letter].width = 18
        for col_letter in ["D", "G", "J"]:
            ws.column_dimensions[col_letter].width = 12

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    filename = f"Timetable_{classes[0]['code']}.xlsx" if len(classes) == 1 else "All_Class_Timetables.xlsx"
    return Response(
        out.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@app.route("/api/export/timetable/pdf", methods=["GET"])
def export_timetable_pdf():
    """
    Exports timetable to a clean, landscape PDF document using reportlab.
    """
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    class_id = request.args.get("class_id", type=int)
    export_all = request.args.get("all", "false").lower() in ("true", "1")

    conn = get_db_connection()
    cursor = conn.cursor()

    if class_id and not export_all:
        cursor.execute("SELECT * FROM classes WHERE id = ?;", (class_id,))
        classes = [dict(r) for r in cursor.fetchall()]
        if not classes:
            conn.close()
            return jsonify({"error": "Class not found."}), 404
    else:
        cursor.execute("SELECT * FROM classes ORDER BY id;")
        classes = [dict(r) for r in cursor.fetchall()]

    cursor.execute("""
        SELECT t.*, 
               s.name AS subject_name, s.code AS subject_code,
               f.name AS faculty_name
        FROM timetable_entries t
        JOIN subjects s ON t.subject_id = s.id
        JOIN faculty f ON t.faculty_id = f.id
        ORDER BY t.class_id, t.day, t.period_num;
    """)
    all_entries = [dict(r) for r in cursor.fetchall()]
    conn.close()

    entries_by_class = {}
    for e in all_entries:
        cid = e["class_id"]
        entries_by_class.setdefault(cid, {}).setdefault(e["day"], {})[e["period_num"]] = e

    pdf_buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        pdf_buffer,
        pagesize=landscape(letter),
        leftMargin=20,
        rightMargin=20,
        topMargin=20,
        bottomMargin=20
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=16,
        alignment=1,
        textColor=colors.HexColor("#0f172a")
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        alignment=1,
        textColor=colors.HexColor("#475569")
    )
    class_title_style = ParagraphStyle(
        "ClassTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        alignment=1,
        textColor=colors.HexColor("#1e3a8a")
    )
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7,
        leading=8.5,
        alignment=1,
        textColor=colors.HexColor("#0f172a")
    )
    header_cell_style = ParagraphStyle(
        "HeaderCell",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9,
        alignment=1,
        textColor=colors.white
    )

    story = []

    headers = [
        "Day / Time",
        f"Period 1\n{PERIOD_TIMINGS[1]['start']}-{PERIOD_TIMINGS[1]['end']}",
        f"Period 2\n{PERIOD_TIMINGS[2]['start']}-{PERIOD_TIMINGS[2]['end']}",
        "Break\n10:40",
        f"Period 3\n{PERIOD_TIMINGS[3]['start']}-{PERIOD_TIMINGS[3]['end']}",
        f"Period 4\n{PERIOD_TIMINGS[4]['start']}-{PERIOD_TIMINGS[4]['end']}",
        "Lunch\n12:30",
        f"Period 5\n{PERIOD_TIMINGS[5]['start']}-{PERIOD_TIMINGS[5]['end']}",
        f"Period 6\n{PERIOD_TIMINGS[6]['start']}-{PERIOD_TIMINGS[6]['end']}",
        "Break\n14:45",
        f"Period 7\n{PERIOD_TIMINGS[7]['start']}-{PERIOD_TIMINGS[7]['end']}",
        f"Period 8\n{PERIOD_TIMINGS[8]['start']}-{PERIOD_TIMINGS[8]['end']}",
    ]

    for idx, c in enumerate(classes):
        if idx > 0:
            story.append(Spacer(1, 25))

        story.append(Paragraph("INFO INSTITUTE OF ENGINEERING (Autonomous)", title_style))
        story.append(Paragraph("Approved by AICTE, New Delhi & Affiliated to Anna University, Chennai", subtitle_style))
        story.append(Paragraph(f"CLASS TIME TABLE: {c['name']} ({c['code']}) — Academic Year 2025-2026", class_title_style))
        story.append(Spacer(1, 8))

        # Build table data
        t_data = []
        # Header
        t_data.append([Paragraph(h.replace("\n", "<br/>"), header_cell_style) for h in headers])

        c_entries = entries_by_class.get(c["id"], {})
        for day in DAYS_OF_WEEK:
            day_grid = c_entries.get(day, {})
            row = [Paragraph(f"<b>{day}</b>", cell_style)]

            for p in [1, 2]:
                e = day_grid.get(p)
                row.append(Paragraph(f"<b>{e['subject_name']}</b><br/>[{e['subject_code']}]" if e else "-", cell_style))

            row.append(Paragraph("<b>BREAK</b>", cell_style))

            for p in [3, 4]:
                e = day_grid.get(p)
                row.append(Paragraph(f"<b>{e['subject_name']}</b><br/>[{e['subject_code']}]" if e else "-", cell_style))

            row.append(Paragraph("<b>LUNCH</b>", cell_style))

            for p in [5, 6]:
                e = day_grid.get(p)
                row.append(Paragraph(f"<b>{e['subject_name']}</b><br/>[{e['subject_code']}]" if e else "-", cell_style))

            row.append(Paragraph("<b>BREAK</b>", cell_style))

            for p in [7, 8]:
                e = day_grid.get(p)
                row.append(Paragraph(f"<b>{e['subject_name']}</b><br/>[{e['subject_code']}]" if e else "-", cell_style))

            t_data.append(row)

        col_widths = [55, 68, 68, 42, 68, 68, 45, 68, 68, 42, 68, 68]
        t = Table(t_data, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ('BACKGROUND', (0, 1), (0, -1), colors.HexColor("#f8fafc")),
            ('BACKGROUND', (3, 1), (3, -1), colors.HexColor("#fef3c7")),
            ('BACKGROUND', (6, 1), (6, -1), colors.HexColor("#fef3c7")),
            ('BACKGROUND', (9, 1), (9, -1), colors.HexColor("#fef3c7")),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t)
        story.append(Spacer(1, 15))

        sig_data = [
            [
                Paragraph("<b>Time Table Coordinator</b>", cell_style),
                Paragraph("<b>HOD / CSE</b>", cell_style),
                Paragraph("<b>Principal</b>", cell_style)
            ]
        ]
        sig_table = Table(sig_data, colWidths=[240, 240, 240])
        sig_table.setStyle(TableStyle([
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        story.append(sig_table)

    doc.build(story)
    pdf_buffer.seek(0)
    filename = f"Timetable_{classes[0]['code']}.pdf" if len(classes) == 1 else "All_Class_Timetables.pdf"
    return Response(
        pdf_buffer.getvalue(),
        mimetype="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@app.route("/api/export/faculty/excel", methods=["GET"])
def export_faculty_excel():
    """
    Exports an individual faculty member's schedule as Microsoft Excel (.xlsx).
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    faculty_id = request.args.get("faculty_id", type=int)
    if not faculty_id:
        return jsonify({"error": "faculty_id parameter is required."}), 400

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM faculty WHERE id = ?;", (faculty_id,))
    fac_row = cursor.fetchone()
    if not fac_row:
        conn.close()
        return jsonify({"error": "Faculty member not found."}), 404
    fac = dict(fac_row)

    cursor.execute("""
        SELECT t.day, t.period_num,
               s.name AS subject_name, s.code AS subject_code,
               c.name AS class_name,  c.code AS class_code
        FROM timetable_entries t
        JOIN subjects s ON t.subject_id = s.id
        JOIN classes  c ON t.class_id   = c.id
        WHERE t.faculty_id = ?
        ORDER BY t.day, t.period_num;
    """, (faculty_id,))
    entries = [dict(r) for r in cursor.fetchall()]
    conn.close()

    grid = {}
    for e in entries:
        grid.setdefault(e["day"], {})[e["period_num"]] = e

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Faculty Schedule"
    ws.sheet_view.showGridLines = True

    header_font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
    sub_font = Font(name="Arial", size=10, italic=True, color="475569")
    table_head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    day_font = Font(name="Arial", size=10, bold=True, color="0F172A")
    cell_font = Font(name="Arial", size=9)
    break_font = Font(name="Arial", size=9, bold=True, color="92400E")

    primary_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    break_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    day_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.merge_cells("A1:L1")
    ws["A1"] = "INFO INSTITUTE OF ENGINEERING (Autonomous)"
    ws["A1"].font = header_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    dept_text = f"Department: {fac['department']}" if fac.get('department') else "Faculty Teaching Schedule"
    ws.merge_cells("A2:L2")
    ws["A2"] = f"FACULTY WEEKLY TEACHING SCHEDULE — {fac['name']} ({dept_text})"
    ws["A2"].font = Font(name="Arial", size=11, bold=True, color="0F172A")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    wtype = "Part-Time (Periods 1-4)" if fac["work_type"] == "part_time" else "Full-Time (Periods 1-8)"
    ws.merge_cells("A3:L3")
    ws["A3"] = f"Work Type: {wtype} | Assigned: {len(entries)} / {fac['weekly_period_limit']} periods/week"
    ws["A3"].font = sub_font
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[5].height = 30

    headers = [
        "Day / Time",
        f"Period 1\n({PERIOD_TIMINGS[1]['start']}-{PERIOD_TIMINGS[1]['end']})",
        f"Period 2\n({PERIOD_TIMINGS[2]['start']}-{PERIOD_TIMINGS[2]['end']})",
        "Morning Break\n(10:40-10:50)",
        f"Period 3\n({PERIOD_TIMINGS[3]['start']}-{PERIOD_TIMINGS[3]['end']})",
        f"Period 4\n({PERIOD_TIMINGS[4]['start']}-{PERIOD_TIMINGS[4]['end']})",
        "Lunch Break\n(12:30-13:05)",
        f"Period 5\n({PERIOD_TIMINGS[5]['start']}-{PERIOD_TIMINGS[5]['end']})",
        f"Period 6\n({PERIOD_TIMINGS[6]['start']}-{PERIOD_TIMINGS[6]['end']})",
        "Tea Break\n(14:45-15:05)",
        f"Period 7\n({PERIOD_TIMINGS[7]['start']}-{PERIOD_TIMINGS[7]['end']})",
        f"Period 8\n({PERIOD_TIMINGS[8]['start']}-{PERIOD_TIMINGS[8]['end']})",
    ]

    for col_idx, col_name in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col_idx, value=col_name)
        cell.font = table_head_font
        cell.fill = primary_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    current_row = 6
    for day in DAYS_OF_WEEK:
        ws.row_dimensions[current_row].height = 34
        day_grid = grid.get(day, {})

        d_cell = ws.cell(row=current_row, column=1, value=day)
        d_cell.font = day_font
        d_cell.fill = day_fill
        d_cell.alignment = Alignment(horizontal="center", vertical="center")
        d_cell.border = thin_border

        for col_idx, p in enumerate([1, 2], start=2):
            e = day_grid.get(p)
            c = ws.cell(row=current_row, column=col_idx, value=f"{e['subject_name']}\n({e['class_code']})" if e else "-")
            c.font = cell_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = thin_border

        mb = ws.cell(row=current_row, column=4, value="BREAK")
        mb.font = break_font
        mb.fill = break_fill
        mb.alignment = Alignment(horizontal="center", vertical="center")
        mb.border = thin_border

        for col_idx, p in enumerate([3, 4], start=5):
            e = day_grid.get(p)
            c = ws.cell(row=current_row, column=col_idx, value=f"{e['subject_name']}\n({e['class_code']})" if e else "-")
            c.font = cell_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = thin_border

        lb = ws.cell(row=current_row, column=7, value="LUNCH")
        lb.font = break_font
        lb.fill = break_fill
        lb.alignment = Alignment(horizontal="center", vertical="center")
        lb.border = thin_border

        for col_idx, p in enumerate([5, 6], start=8):
            e = day_grid.get(p)
            c = ws.cell(row=current_row, column=col_idx, value=f"{e['subject_name']}\n({e['class_code']})" if e else "-")
            c.font = cell_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = thin_border

        tb = ws.cell(row=current_row, column=10, value="BREAK")
        tb.font = break_font
        tb.fill = break_fill
        tb.alignment = Alignment(horizontal="center", vertical="center")
        tb.border = thin_border

        for col_idx, p in enumerate([7, 8], start=11):
            e = day_grid.get(p)
            c = ws.cell(row=current_row, column=col_idx, value=f"{e['subject_name']}\n({e['class_code']})" if e else "-")
            c.font = cell_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = thin_border

        current_row += 1

    sig_row = current_row + 2
    ws.cell(row=sig_row, column=2, value="Faculty Signature").font = Font(name="Arial", size=10, bold=True)
    ws.cell(row=sig_row, column=6, value="HOD / CSE").font = Font(name="Arial", size=10, bold=True)
    ws.cell(row=sig_row, column=11, value="Principal").font = Font(name="Arial", size=10, bold=True)

    ws.column_dimensions["A"].width = 14
    for col_letter in ["B", "C", "E", "F", "H", "I", "K", "L"]:
        ws.column_dimensions[col_letter].width = 18
    for col_letter in ["D", "G", "J"]:
        ws.column_dimensions[col_letter].width = 12

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    safe_name = "".join(c for c in fac["name"] if c.isalnum() or c in (' ', '_')).rstrip().replace(' ', '_')
    return Response(
        out.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Faculty_Timetable_{safe_name}.xlsx"}
    )


@app.route("/api/export/faculty/pdf", methods=["GET"])
def export_faculty_pdf():
    """
    Exports an individual faculty member's schedule as landscape PDF using reportlab.
    """
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    faculty_id = request.args.get("faculty_id", type=int)
    if not faculty_id:
        return jsonify({"error": "faculty_id parameter is required."}), 400

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM faculty WHERE id = ?;", (faculty_id,))
    fac_row = cursor.fetchone()
    if not fac_row:
        conn.close()
        return jsonify({"error": "Faculty member not found."}), 404
    fac = dict(fac_row)

    cursor.execute("""
        SELECT t.day, t.period_num,
               s.name AS subject_name, s.code AS subject_code,
               c.name AS class_name,  c.code AS class_code
        FROM timetable_entries t
        JOIN subjects s ON t.subject_id = s.id
        JOIN classes  c ON t.class_id   = c.id
        WHERE t.faculty_id = ?
        ORDER BY t.day, t.period_num;
    """, (faculty_id,))
    entries = [dict(r) for r in cursor.fetchall()]
    conn.close()

    grid = {}
    for e in entries:
        grid.setdefault(e["day"], {})[e["period_num"]] = e

    pdf_buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        pdf_buffer,
        pagesize=landscape(letter),
        leftMargin=20,
        rightMargin=20,
        topMargin=20,
        bottomMargin=20
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=16,
        alignment=1,
        textColor=colors.HexColor("#0f172a")
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        alignment=1,
        textColor=colors.HexColor("#475569")
    )
    faculty_title_style = ParagraphStyle(
        "FacultyTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        alignment=1,
        textColor=colors.HexColor("#1e3a8a")
    )
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7,
        leading=8.5,
        alignment=1,
        textColor=colors.HexColor("#0f172a")
    )
    header_cell_style = ParagraphStyle(
        "HeaderCell",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9,
        alignment=1,
        textColor=colors.white
    )

    story = []
    story.append(Paragraph("INFO INSTITUTE OF ENGINEERING (Autonomous)", title_style))
    story.append(Paragraph("Approved by AICTE, New Delhi & Affiliated to Anna University, Chennai", subtitle_style))

    dept_text = f"Department: {fac['department']}" if fac.get('department') else "Department Faculty"
    wtype = "Part-Time (Periods 1-4)" if fac["work_type"] == "part_time" else "Full-Time (Periods 1-8)"
    story.append(Paragraph(f"FACULTY SCHEDULE: {fac['name']} — {dept_text} | {wtype} | Assigned: {len(entries)} / {fac['weekly_period_limit']} periods/week", faculty_title_style))
    story.append(Spacer(1, 10))

    headers = [
        "Day / Time",
        f"Period 1\n{PERIOD_TIMINGS[1]['start']}-{PERIOD_TIMINGS[1]['end']}",
        f"Period 2\n{PERIOD_TIMINGS[2]['start']}-{PERIOD_TIMINGS[2]['end']}",
        "Break\n10:40",
        f"Period 3\n{PERIOD_TIMINGS[3]['start']}-{PERIOD_TIMINGS[3]['end']}",
        f"Period 4\n{PERIOD_TIMINGS[4]['start']}-{PERIOD_TIMINGS[4]['end']}",
        "Lunch\n12:30",
        f"Period 5\n{PERIOD_TIMINGS[5]['start']}-{PERIOD_TIMINGS[5]['end']}",
        f"Period 6\n{PERIOD_TIMINGS[6]['start']}-{PERIOD_TIMINGS[6]['end']}",
        "Break\n14:45",
        f"Period 7\n{PERIOD_TIMINGS[7]['start']}-{PERIOD_TIMINGS[7]['end']}",
        f"Period 8\n{PERIOD_TIMINGS[8]['start']}-{PERIOD_TIMINGS[8]['end']}",
    ]

    t_data = []
    t_data.append([Paragraph(h.replace("\n", "<br/>"), header_cell_style) for h in headers])

    for day in DAYS_OF_WEEK:
        day_grid = grid.get(day, {})
        row = [Paragraph(f"<b>{day}</b>", cell_style)]

        for p in [1, 2]:
            e = day_grid.get(p)
            row.append(Paragraph(f"<b>{e['subject_name']}</b><br/>({e['class_code']})" if e else "-", cell_style))

        row.append(Paragraph("<b>BREAK</b>", cell_style))

        for p in [3, 4]:
            e = day_grid.get(p)
            row.append(Paragraph(f"<b>{e['subject_name']}</b><br/>({e['class_code']})" if e else "-", cell_style))

        row.append(Paragraph("<b>LUNCH</b>", cell_style))

        for p in [5, 6]:
            e = day_grid.get(p)
            row.append(Paragraph(f"<b>{e['subject_name']}</b><br/>({e['class_code']})" if e else "-", cell_style))

        row.append(Paragraph("<b>BREAK</b>", cell_style))

        for p in [7, 8]:
            e = day_grid.get(p)
            row.append(Paragraph(f"<b>{e['subject_name']}</b><br/>({e['class_code']})" if e else "-", cell_style))

        t_data.append(row)

    col_widths = [55, 68, 68, 42, 68, 68, 45, 68, 68, 42, 68, 68]
    t = Table(t_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ('BACKGROUND', (0, 1), (0, -1), colors.HexColor("#f8fafc")),
        ('BACKGROUND', (3, 1), (3, -1), colors.HexColor("#fef3c7")),
        ('BACKGROUND', (6, 1), (6, -1), colors.HexColor("#fef3c7")),
        ('BACKGROUND', (9, 1), (9, -1), colors.HexColor("#fef3c7")),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t)
    story.append(Spacer(1, 20))

    sig_data = [
        [
            Paragraph("<b>Faculty Signature</b>", cell_style),
            Paragraph("<b>HOD / CSE</b>", cell_style),
            Paragraph("<b>Principal</b>", cell_style)
        ]
    ]
    sig_table = Table(sig_data, colWidths=[240, 240, 240])
    sig_table.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(sig_table)

    doc.build(story)
    pdf_buffer.seek(0)
    safe_name = "".join(c for c in fac["name"] if c.isalnum() or c in (' ', '_')).rstrip().replace(' ', '_')
    return Response(
        pdf_buffer.getvalue(),
        mimetype="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=Faculty_Timetable_{safe_name}.pdf"}
    )


@app.route("/api/export/faculty/csv", methods=["GET"])
def export_faculty_csv():
    """
    Exports an individual faculty member's weekly schedule as a CSV spreadsheet.
    """
    faculty_id = request.args.get("faculty_id", type=int)
    if not faculty_id:
        return jsonify({"error": "faculty_id parameter is required."}), 400

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM faculty WHERE id = ?;", (faculty_id,))
    fac_row = cursor.fetchone()
    if not fac_row:
        conn.close()
        return jsonify({"error": "Faculty member not found."}), 404
    fac = dict(fac_row)

    cursor.execute("""
        SELECT t.day, t.period_num,
               s.name AS subject_name, s.code AS subject_code,
               c.name AS class_name,  c.code AS class_code
        FROM timetable_entries t
        JOIN subjects s ON t.subject_id = s.id
        JOIN classes  c ON t.class_id   = c.id
        WHERE t.faculty_id = ?
        ORDER BY t.day, t.period_num;
    """, (faculty_id,))
    entries = [dict(r) for r in cursor.fetchall()]
    conn.close()

    grid = {}
    for e in entries:
        grid.setdefault(e["day"], {})[e["period_num"]] = e

    output = io.StringIO()
    output.write("\ufeff")  # UTF-8 BOM
    writer = csv.writer(output)

    writer.writerow(["INFO INSTITUTE OF ENGINEERING (Autonomous)"])
    writer.writerow(["FACULTY WEEKLY TEACHING SCHEDULE & LOAD ALLOCATION"])
    writer.writerow([f"Faculty Name: {fac['name']}", f"Department: {fac.get('department') or ''}"])
    wtype = "Part-Time (09:00 AM - 12:30 PM, Periods 1-4)" if fac["work_type"] == "part_time" else "Full-Time (09:00 AM - 04:45 PM, All 8 Periods)"
    writer.writerow([f"Work Type: {wtype}", f"Assigned Periods: {len(entries)} / {fac['weekly_period_limit']} periods/week"])
    writer.writerow([])

    headers = [
        "Day / Time",
        f"Period 1 ({PERIOD_TIMINGS[1]['start']} - {PERIOD_TIMINGS[1]['end']})",
        f"Period 2 ({PERIOD_TIMINGS[2]['start']} - {PERIOD_TIMINGS[2]['end']})",
        "Morning Break (10:40 - 10:50)",
        f"Period 3 ({PERIOD_TIMINGS[3]['start']} - {PERIOD_TIMINGS[3]['end']})",
        f"Period 4 ({PERIOD_TIMINGS[4]['start']} - {PERIOD_TIMINGS[4]['end']})",
        "Lunch Break (12:30 - 13:05)",
        f"Period 5 ({PERIOD_TIMINGS[5]['start']} - {PERIOD_TIMINGS[5]['end']})",
        f"Period 6 ({PERIOD_TIMINGS[6]['start']} - {PERIOD_TIMINGS[6]['end']})",
        "Tea Break (14:45 - 15:05)",
        f"Period 7 ({PERIOD_TIMINGS[7]['start']} - {PERIOD_TIMINGS[7]['end']})",
        f"Period 8 ({PERIOD_TIMINGS[8]['start']} - {PERIOD_TIMINGS[8]['end']})",
    ]
    writer.writerow(headers)

    for day in DAYS_OF_WEEK:
        day_grid = grid.get(day, {})
        row = [day]
        for p in range(1, 9):
            if p == 3:
                row.append("MORNING BREAK")
            elif p == 5:
                row.append("LUNCH BREAK")
            elif p == 7:
                row.append("TEA BREAK")
            e = day_grid.get(p)
            row.append(f"{e['subject_name']} - {e['class_code']}" if e else "-")
        writer.writerow(row)

    writer.writerow([])
    writer.writerow(["Faculty Signature", "", "", "HOD / CSE", "", "", "", "Principal"])

    safe_name = "".join(c for c in fac["name"] if c.isalnum() or c in (' ', '_')).rstrip().replace(' ', '_')
    return Response(
        output.getvalue(),
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename=Faculty_Timetable_{safe_name}.csv"}
    )


@app.route("/api/timetable/regenerate", methods=["POST"])
def manual_regenerate_timetable():
    """
    Manually triggers timetable regeneration and returns detailed diagnostics.
    """
    try:
        scheduler = TimetableScheduler()
        result = scheduler.generate()
        return jsonify(result)
    except SchedulingError as se:
        return jsonify({
            "success": False,
            "error": se.message,
            "details": se.details,
        }), 422
    except Exception as e:
        return jsonify({
            "success": False,
            "error": "An unexpected error occurred during timetable generation.",
        }), 500


# ---------------------------------------------------------------------------
# Supabase Auth Alias Endpoints (alias routes for backwards compatibility)
# ---------------------------------------------------------------------------
@app.route("/api/auth/signup", methods=["POST"])
def supabase_signup():
    """Alias for /api/signup — forwards to the primary handler."""
    return signup()


@app.route("/api/auth/login", methods=["POST"])
def supabase_login():
    """Alias for /api/login — forwards to the primary handler."""
    return login()


@app.route("/api/auth/logout", methods=["POST"])
def supabase_logout():
    """Signs out current user session."""
    try:
        auth_header = request.headers.get("Authorization", "")
        if auth_header.startswith("Bearer "):
            token = auth_header.split(" ")[1]
            supabase.auth.admin.sign_out(token)
        else:
            supabase.auth.sign_out()
        return jsonify({"success": True, "message": "Logged out successfully."})
    except Exception as e:
        return jsonify({"success": True, "message": "Session ended."})


@app.route("/api/auth/user", methods=["GET"])
def supabase_get_user():
    """Returns the authenticated user details from Supabase token."""
    auth_header = request.headers.get("Authorization", "")
    if not auth_header.startswith("Bearer "):
        return jsonify({"error": "Missing or invalid Bearer token."}), 401

    token = auth_header.split(" ")[1]
    try:
        user_res = supabase.auth.get_user(token)
        if not user_res or not user_res.user:
            return jsonify({"error": "Invalid session token."}), 401
        u = user_res.user
        return jsonify({
            "success": True,
            "user": {
                "id": u.id,
                "email": u.email,
                "created_at": getattr(u, "created_at", None),
            }
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 401


@app.route("/api/supabase/status", methods=["GET"])
def supabase_status():
    """Health check for Supabase client connection and table status."""
    try:
        tables = ["classes", "faculty", "subjects", "timetable_entries"]
        status = {}
        for tbl in tables:
            r = supabase.table(tbl).select("id", count="exact").limit(1).execute()
            status[tbl] = {"available": True, "count": r.count}
        return jsonify({
            "success": True,
            "supabase_url": SUPABASE_URL,
            "connected": True,
            "tables": status
        })
    except Exception as e:
        return jsonify({
            "success": False,
            "connected": False,
            "error": str(e)
        }), 500


# ---------------------------------------------------------------------------
# Internal Examination Seating & Hall Allocation APIs (Supabase)
# ---------------------------------------------------------------------------
@app.route("/api/exam/halls", methods=["GET", "POST"])
def exam_halls():
    """
    GET: Lists all examination halls.
    POST: Inserts a new examination hall (hall_no, capacity).
    """
    if request.method == "POST":
        data = request.get_json() or {}
        hall_no = (data.get("hall_no") or "").strip().upper()
        capacity_raw = data.get("capacity")

        if not hall_no:
            return jsonify({"error": "Hall number is required."}), 400
        try:
            capacity = int(capacity_raw) if capacity_raw is not None else 40
            if capacity < 1:
                return jsonify({"error": "Capacity must be at least 1."}), 400
        except (ValueError, TypeError):
            return jsonify({"error": "Invalid capacity value."}), 400

        try:
            insert_payload = {"hall_no": hall_no, "capacity": capacity}
            if data.get("building"):
                insert_payload["building"] = data.get("building")
            res = supabase.table("exam_halls").insert(insert_payload).execute()
            return jsonify({
                "success": True,
                "message": f"Exam hall '{hall_no}' added successfully.",
                "hall": res.data[0] if res.data else insert_payload
            }), 201
        except Exception as e:
            return jsonify({"error": str(e)}), 400

    # GET request
    try:
        res = supabase.table("exam_halls").select("*").order("hall_no").execute()
        return jsonify(res.data or [])
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/exam/allocations", methods=["GET", "POST"])
def exam_allocations():
    """
    GET: Returns complete seating allocations joined across exam_allocations,
         exam_halls, classes, exam_schedules, and subjects in a single JSON payload.
    POST: Saves a new exam hall seating arrangement.
    """
    if request.method == "POST":
        data = request.get_json() or {}
        hall_id = data.get("hall_id")
        class_id = data.get("class_id")
        exam_schedule_id = data.get("exam_schedule_id")
        roll_no_start = (data.get("roll_no_start") or data.get("roll_start") or "").strip()
        roll_no_end = (data.get("roll_no_end") or data.get("roll_end") or "").strip()
        student_count = data.get("student_count") or data.get("total_students")

        if not hall_id or not class_id or not exam_schedule_id:
            return jsonify({"error": "hall_id, class_id, and exam_schedule_id are required."}), 400
        if not roll_no_start or not roll_no_end:
            return jsonify({"error": "Roll number start and end are required."}), 400

        try:
            count_val = int(student_count) if student_count is not None else 0
        except (ValueError, TypeError):
            count_val = 0

        payload = {
            "hall_id": hall_id,
            "class_id": class_id,
            "exam_schedule_id": exam_schedule_id,
            "roll_no_start": roll_no_start,
            "roll_no_end": roll_no_end,
            "student_count": count_val
        }

        try:
            res = supabase.table("exam_allocations").insert(payload).execute()
            return jsonify({
                "success": True,
                "message": "Exam allocation saved successfully.",
                "allocation": res.data[0] if res.data else payload
            }), 201
        except Exception as e:
            return jsonify({"error": str(e)}), 400

    # GET: Join query via Supabase PostgREST
    try:
        res = supabase.table("exam_allocations").select("""
            id,
            roll_no_start,
            roll_no_end,
            student_count,
            created_at,
            exam_halls ( id, hall_no, capacity ),
            classes ( id, name, code ),
            exam_schedules ( id, exam_name, exam_date, session, subjects ( id, name, code ) )
        """).order("id").execute()

        formatted_allocations = []
        for item in (res.data or []):
            hall_info = item.get("exam_halls") or {}
            class_info = item.get("classes") or {}
            sched_info = item.get("exam_schedules") or {}
            subj_info = sched_info.get("subjects") or {}

            # Construct roll number range
            start_roll = item.get("roll_no_start") or ""
            end_roll = item.get("roll_no_end") or ""
            roll_range = f"{start_roll} - {end_roll}" if start_roll and end_roll else (start_roll or end_roll or "N/A")

            # Format Date & Session
            date_str = sched_info.get("exam_date") or ""
            session_str = sched_info.get("session") or ""
            date_session = f"{date_str} ({session_str})" if date_str and session_str else (date_str or session_str or "N/A")

            formatted_allocations.append({
                "id": item.get("id"),
                "hall_no": hall_info.get("hall_no", "N/A"),
                "hall_capacity": hall_info.get("capacity"),
                "date_session": date_session,
                "exam_date": date_str,
                "session": session_str,
                "subject_name": subj_info.get("name", "N/A"),
                "subject_code": subj_info.get("code", ""),
                "class_name": class_info.get("name", "N/A"),
                "class_code": class_info.get("code", ""),
                "roll_number_range": roll_range,
                "roll_no_start": start_roll,
                "roll_no_end": end_roll,
                "total_students": item.get("student_count", 0),
                "exam_name": sched_info.get("exam_name", "Internal Examination")
            })

        return jsonify(formatted_allocations)
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ---------------------------------------------------------------------------
# Global Error Handlers (Clean JSON responses, no tracebacks exposed)
# ---------------------------------------------------------------------------
@app.errorhandler(404)
def not_found_handler(e):
    return jsonify({"error": "The requested resource was not found."}), 404


@app.errorhandler(405)
def method_not_allowed_handler(e):
    return jsonify({"error": "HTTP method not allowed on this endpoint."}), 405


@app.errorhandler(500)
def server_error_handler(e):
    return jsonify({"error": "An internal server error occurred. Please verify your inputs."}), 500


if __name__ == "__main__":
    # Local development server
    init_db()
    seed_sample_data_if_empty()
    try:
        TimetableScheduler().generate()
    except Exception:
        pass
    print("Intelligent Department Timetable Scheduler running at http://127.0.0.1:5000")
    app.run(host="127.0.0.1", port=5000, debug=True)
